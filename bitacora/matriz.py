from datetime import date, datetime
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from tickets.models import Ticket
from .models import Actividad
COLOR_ENCABEZADO = '1F3A5F'
BORDE = Border(*(Side(style='thin', color='D0D5DD'),) * 4)

def _local(valor):
    if isinstance(valor, str) and valor[:1] in ('=', '+', '-', '@', chr(9), chr(13), chr(10)):
        return "'" + valor
    if isinstance(valor, datetime) and timezone.is_aware(valor):
        return timezone.localtime(valor).replace(tzinfo=None, microsecond=0)
    return valor

def _horas(delta):
    return round(delta.total_seconds() / 3600, 2) if delta is not None else ''

def _hoja(libro, titulo, encabezados, filas, anchos):
    hoja = libro.create_sheet(titulo)
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = BORDE
    for fila in filas:
        hoja.append([_local(v) for v in fila])
    hoja.freeze_panes = 'A2'
    if hoja.max_row > 1:
        hoja.auto_filter.ref = f'A1:{get_column_letter(len(encabezados))}{hoja.max_row}'
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.border = BORDE
            celda.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(celda.value, datetime):
                celda.number_format = 'yyyy-mm-dd hh:mm'
            elif isinstance(celda.value, date):
                celda.number_format = 'yyyy-mm-dd'
    return hoja

def construir_matriz(tecnico=None, desde=None, hasta=None):
    actividades = Actividad.objects.select_related('tecnico', 'equipo', 'ticket').order_by('fecha', 'id')
    tickets = Ticket.objects.select_related('solicitante', 'tecnico', 'categoria', 'equipo').order_by('id')
    if tecnico is not None:
        actividades = actividades.filter(tecnico=tecnico)
        tickets = tickets.filter(tecnico=tecnico)
    else:
        tickets = tickets.filter(tecnico__isnull=False)
    if desde:
        actividades = actividades.filter(fecha__gte=desde)
        tickets = tickets.filter(creado__date__gte=desde)
    if hasta:
        actividades = actividades.filter(fecha__lte=hasta)
        tickets = tickets.filter(creado__date__lte=hasta)
    libro = Workbook()
    libro.remove(libro.active)
    filas_a = [(a.fecha, a.tecnico.nombre_completo, a.get_tipo_display(), a.titulo, a.descripcion, a.duracion_minutos, a.equipo.codigo if a.equipo else '', a.ticket.codigo if a.ticket else '', 'Automática (ticket)' if a.automatica else 'Manual') for a in actividades]
    _hoja(libro, 'Actividades', ['Fecha', 'Técnico', 'Tipo', 'Actividad', 'Detalle', 'Minutos', 'Equipo', 'Ticket', 'Origen'], filas_a, [12, 24, 26, 36, 60, 9, 14, 10, 18])
    filas_t, resueltos, cumplidos = ([], 0, 0)
    for t in tickets:
        cumplio = t.cumplio_sla
        if t.fecha_resolucion:
            resueltos += 1
            cumplidos += 1 if cumplio else 0
        filas_t.append((t.codigo, t.creado, t.solicitante.nombre_completo, t.solicitante.departamento, t.titulo, t.categoria.nombre, t.get_prioridad_display(), t.get_estado_display(), t.tecnico.nombre_completo, t.equipo.codigo if t.equipo else '', t.fecha_primera_atencion, t.fecha_resolucion, t.fecha_cierre, t.fecha_limite, _horas(t.tiempo_resolucion), '' if cumplio is None else 'Sí' if cumplio else 'No', t.reprogramaciones, t.solucion))
    _hoja(libro, 'Tickets atendidos', ['Ticket', 'Creado', 'Solicitante', 'Departamento', 'Asunto', 'Categoría', 'Prioridad', 'Estado', 'Técnico', 'Equipo', '1ª atención', 'Resuelto', 'Cerrado', 'Vence (SLA)', 'Horas hasta resolver', '¿Cumplió SLA?', 'Veces reprogramado', 'Solución'], filas_t, [10, 17, 24, 18, 34, 18, 11, 14, 22, 14, 17, 17, 17, 17, 12, 10, 12, 60])
    total_min = sum((f[5] for f in filas_a))
    resumen = libro.create_sheet('Resumen', 0)
    datos = [('MATRIZ DE ACTIVIDADES - ÁREA DE TICS', ''), ('Técnico', tecnico.nombre_completo if tecnico else 'Todos los técnicos'), ('Generada el', _local(timezone.now())), ('Período', f"{desde or 'inicio'}  a  {hasta or 'hoy'}"), ('', ''), ('Actividades registradas', len(filas_a)), ('Horas registradas', round(total_min / 60, 1)), ('Tickets atendidos', len(filas_t)), ('Tickets resueltos', resueltos), ('Cumplimiento de SLA', f'{cumplidos * 100 // resueltos} %' if resueltos else 'n/a')]
    for fila in datos:
        resumen.append([_local(v) for v in fila])
    resumen['A1'].font = Font(bold=True, size=14, color=COLOR_ENCABEZADO)
    for celda in resumen['A'][1:]:
        celda.font = Font(bold=True)
    resumen.column_dimensions['A'].width = 28
    resumen.column_dimensions['B'].width = 50
    resumen['B3'].number_format = 'yyyy-mm-dd hh:mm'
    for fila in resumen.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].alignment = Alignment(horizontal='left')
    return libro
